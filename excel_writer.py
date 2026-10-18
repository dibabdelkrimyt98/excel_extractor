
"""Écriture du classeur de sortie et du journal de traitement."""
from __future__ import annotations

import time
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def copy_cell_format(source: Cell, target: Cell) -> None:
    """Copie le format et les propriétés utiles d'une cellule source."""
    if source.has_style:
        target._style = copy(source._style)

    if source.number_format:
        target.number_format = source.number_format

    if source.alignment:
        target.alignment = copy(source.alignment)

    if source.protection:
        target.protection = copy(source.protection)

    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)

    if source.comment:
        target.comment = copy(source.comment)


def write_result(data: dict[str, Any], config: dict[str, Any]) -> Path:
    """Écrit le fichier résultat et son journal, puis retourne son chemin."""
    write_start = time.perf_counter()

    output_path: Path = data["output_path"]

    if output_path.exists() and not config["overwrite_output"]:
        raise FileExistsError(
            "Le fichier de sortie existe déjà et overwrite_output=false : "
            f"{output_path}"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)

    journal_path = output_path.with_name(
        f"{output_path.stem}_journal.txt"
    )

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Resultats"

    output_fields = data["output_fields"]

    # En-têtes.
    for col_idx, field_name in enumerate(output_fields, start=1):
        cell = ws_out.cell(
            row=1,
            column=col_idx,
            value=field_name,
        )
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws_out.freeze_panes = "A2"
    ws_out.auto_filter.ref = (
        f"A1:{get_column_letter(len(output_fields))}1"
    )

    changes_by_key = {
        (
            item["file_key"],
            item["source_row"],
            item["field_name"],
        ): item
        for item in data["change_log"]
    }

    try:
        with journal_path.open("w", encoding="utf-8") as log_file:
            log_file.write(
                "Journal Excel Extractor - fusion de deux fichiers\n"
            )
            log_file.write(f"Fichier de sortie : {output_path}\n")
            log_file.write(
                f"Nombre de lignes exportées : {len(data['records'])}\n"
            )
            log_file.write(
                "Nombre de cellules remplacées : "
                f"{len(data['change_log'])}\n"
            )
            log_file.write(
                "Nombre de lignes exclues pour motif : "
                f"{data.get('excluded_motif_count', 0)}\n"
            )
            log_file.write(
                "Nombre de lignes exclues pour date : "
                f"{data.get('excluded_date_count', 0)}\n"
            )
            log_file.write("\n")

            # Événements et avertissements enregistrés pendant la lecture.
            for warning in data.get("warnings", []):
                message = warning.get("message", str(warning))
                log_file.write(message + "\n")

            log_file.write("\n--- Remplacements de cellules ---\n")

            for out_row, record in enumerate(data["records"], start=2):
                for out_col, field_name in enumerate(
                    output_fields,
                    start=1,
                ):
                    target_coordinate = (
                        f"{get_column_letter(out_col)}{out_row}"
                    )

                    if field_name not in record["values_by_field"]:
                        # Champ absent du schéma de ce fichier source.
                        ws_out.cell(
                            row=out_row,
                            column=out_col,
                            value=config["empty_value"],
                        )
                        continue

                    value = record["values_by_field"][field_name]

                    target = ws_out.cell(
                        row=out_row,
                        column=out_col,
                        value=value,
                    )

                    source_cell = record["styles_by_field"][field_name]
                    copy_cell_format(source_cell, target)

                    change_key = (
                        record["file_key"],
                        record["source_row"],
                        field_name,
                    )
                    change = changes_by_key.get(change_key)

                    if change:
                        source_coordinate = change["source_coordinate"]

                        message = (
                            f"[REMPLACEMENT][ÉCRITURE] "
                            f"fichier={record['source_file']} | "
                            f"source={record['source_sheet']}!"
                            f"{source_coordinate} "
                            f"(ligne={record['source_row']}, "
                            f"colonne={change['source_column']}) "
                            f"-> sortie=Resultats!{target_coordinate} "
                            f"(ligne={out_row}, "
                            f"colonne={get_column_letter(out_col)}) "
                            f"| champ={field_name!r} "
                            f"| valeur={value!r}"
                        )

                        print(message, flush=True)
                        log_file.write(message + "\n")

            # Ajustement de largeur à partir des en-têtes et des
            # 100 premières lignes de données, pour limiter le coût.
            for col_idx, field_name in enumerate(
                output_fields,
                start=1,
            ):
                max_length = len(str(field_name))

                last_width_row = min(ws_out.max_row, 101)

                for row_idx in range(2, last_width_row + 1):
                    value = ws_out.cell(
                        row=row_idx,
                        column=col_idx,
                    ).value

                    if value is not None:
                        max_length = max(
                            max_length,
                            len(str(value)),
                        )

                ws_out.column_dimensions[
                    get_column_letter(col_idx)
                ].width = min(max(max_length + 2, 12), 32)

            # Mesure incluant la préparation et la sauvegarde du résultat.
            wb_out.save(output_path)

        write_seconds = time.perf_counter() - write_start
        data["write_seconds"] = write_seconds

        timing_message = (
            f"[TEMPS D'ÉCRITURE] Génération et sauvegarde de "
            f"{output_path.name} : {write_seconds:.3f} seconde(s)"
        )

        print(timing_message, flush=True)

        # Ajouter la durée au journal après la sauvegarde.
        with journal_path.open("a", encoding="utf-8") as log_file:
            log_file.write("\n" + timing_message + "\n")

        print(f"[JOURNAL] {journal_path}", flush=True)

        return output_path

    except Exception as exc:
        message = (
            f"[ERREUR][ÉCRITURE/SAUVEGARDE] "
            f"{type(exc).__name__}: {exc}"
        )

        print(message, flush=True)

        try:
            with journal_path.open("a", encoding="utf-8") as log_file:
                log_file.write(message + "\n")
        except OSError:
            pass

        raise

    finally:
        wb_out.close()
