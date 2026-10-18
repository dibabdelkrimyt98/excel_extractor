"""Création du classeur de statistiques et des graphiques."""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DARK_BLUE = "17365D"
WHITE = "FFFFFF"


def style_header(worksheet, row=1):
    for cell in worksheet[row]:
        if cell.value is not None:
            cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def autofit_columns(worksheet, max_width=35):
    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), max_width)


def apply_borders(worksheet):
    side = Side(style="thin", color="D9E1F2")
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_statistics(gender_stats, age_stats, output_file, employees_data):
    output_path = Path(output_file).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard["A1"] = "STATISTIQUES DES EMPLOYÉS"
    dashboard["A1"].font = Font(size=16, bold=True, color=DARK_BLUE)
    dashboard.merge_cells("A1:B1")
    summary = [
        ("Indicateur", "Valeur"),
        ("Total employés", gender_stats["total"]),
        ("Femmes", gender_stats["counts"]["Femme"]),
        ("Hommes", gender_stats["counts"]["Homme"]),
        ("Civilité VIDE / inconnue", gender_stats["counts"]["VIDE"]),
        ("Dates de naissance absentes ou invalides", age_stats["invalid_birth_dates"]),
        ("Dates de naissance postérieures à la référence", age_stats["future_birth_dates"]),
        ("Date de référence", age_stats["reference_date"]),
        ("Lignes exclues (Code Depart = 18)", age_stats["code_depart_filtered_count"]),
        ("Lignes vides ignorées", age_stats["empty_rows_skipped"]),
    ]
    for row in summary:
        dashboard.append(row)
    style_header(dashboard, 2)

    pie = PieChart()
    pie.title = "Répartition par sexe"
    pie.height, pie.width = 8, 12
    pie.add_data(Reference(dashboard, min_col=2, min_row=4, max_row=6), titles_from_data=False)
    pie.set_categories(Reference(dashboard, min_col=1, min_row=4, max_row=6))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showLeaderLines = True
    dashboard.add_chart(pie, "D2")

    bands_sheet = workbook.create_sheet("Tranches d'âge")
    bands_sheet.append(["Tranche d'âge", "Femme", "Homme", "VIDE", "Total"])
    for label, counts in age_stats["bands"].items():
        bands_sheet.append([label, counts["Femme"], counts["Homme"], counts["VIDE"], sum(counts.values())])
    style_header(bands_sheet, 1)

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Effectifs par tranche d'âge et sexe"
    chart.y_axis.title = "Nombre d'employés"
    chart.x_axis.title = "Tranche d'âge"
    chart.height, chart.width = 9, 17
    chart.add_data(Reference(bands_sheet, min_col=2, max_col=4, min_row=1, max_row=bands_sheet.max_row), titles_from_data=True)
    chart.set_categories(Reference(bands_sheet, min_col=1, min_row=2, max_row=bands_sheet.max_row))
    bands_sheet.add_chart(chart, "G2")

    # Tableau des âges exacts : écrire aux lignes précises, sans append décalé.
    exact_start = bands_sheet.max_row + 3
    bands_sheet.cell(exact_start, 1, "Âge révolu")
    bands_sheet.cell(exact_start, 2, "Nombre d'employés")
    for offset, (age, count) in enumerate(age_stats["exact_ages"].items(), start=1):
        bands_sheet.cell(exact_start + offset, 1, age)
        bands_sheet.cell(exact_start + offset, 2, count)
    style_header(bands_sheet, exact_start)

    detail = workbook.create_sheet("Détail employés")
    config = employees_data["config"]
    source_columns = list(config["columns"].keys())
    detail.append(source_columns + ["Sexe calculé", "Âge révolu au " + age_stats["reference_date"], "Tranche d'âge"])
    for employee in employees_data["employees"]:
        row = [employee["values"].get(column, config.get("empty_value", "VIDE")) for column in source_columns]
        for index, value in enumerate(row):
            if value is None or (isinstance(value, str) and not value.strip()):
                row[index] = config.get("empty_value", "VIDE")
        row.extend([employee["gender"], employee["age"] if employee["age"] is not None else config.get("empty_value", "VIDE"), employee["age_band"]])
        detail.append(row)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        style_header(worksheet, 1)
        autofit_columns(worksheet)
        apply_borders(worksheet)
        worksheet.sheet_view.showGridLines = False
    style_header(dashboard, 2)
    workbook.save(output_path)
    workbook.close()
    return output_path
