
"""Fonctions communes de lecture, validation et préparation des statistiques Excel."""

from __future__ import annotations

import json
import math
import shutil
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils.cell import column_index_from_string
from openpyxl.utils.datetime import from_excel


DEFAULT_DATE_FORMATS = ("%m/%d/%Y",)
SUPPORTED_EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}

REQUIRED_CONFIG_KEYS = (
    "input_file",
    "output_file",
    "sheet_name",
    "start_row",
    "reference_date",
    "columns",
    "age_bands",
)

REQUIRED_EMPLOYEE_COLUMNS = {"CIVILITE", "DATE NAISSANCE"}


# ============================================================
# CONFIGURATION
# ============================================================

def _validate_excel_column(name: str, value: Any) -> str:
    """Valide une lettre de colonne Excel et retourne sa forme majuscule."""

    if not isinstance(value, str) or not value.strip():
        raise ValueError(
            f"La colonne configurée pour {name!r} doit être "
            "une lettre Excel, par exemple 'A' ou 'AA'."
        )

    letter = value.strip().upper()

    try:
        column_index_from_string(letter)
    except ValueError as exc:
        raise ValueError(
            f"Colonne Excel invalide pour {name!r} : {value!r}"
        ) from exc

    return letter


def _validate_age_bands(age_bands: Any) -> list[dict[str, Any]]:
    """Valide les tranches d'âge et vérifie qu'elles ne se chevauchent pas."""

    if not isinstance(age_bands, list) or not age_bands:
        raise ValueError("'age_bands' doit être une liste non vide.")

    validated = []
    labels = set()

    for band in age_bands:
        if not isinstance(band, dict):
            raise ValueError("Chaque tranche d'âge doit être un objet JSON.")

        if not {"label", "min", "max"} <= set(band):
            raise ValueError(
                "Chaque tranche doit contenir 'label', 'min' et 'max'."
            )

        label = band["label"]
        minimum = band["min"]
        maximum = band["max"]

        if not isinstance(label, str) or not label.strip():
            raise ValueError("Chaque tranche doit avoir un label non vide.")

        if label in labels:
            raise ValueError(f"Label de tranche en double : {label!r}")

        if (
            isinstance(minimum, bool)
            or not isinstance(minimum, int)
            or minimum < 0
        ):
            raise ValueError(f"Borne minimale invalide pour {label!r}.")

        if maximum is not None and (
            isinstance(maximum, bool)
            or not isinstance(maximum, int)
            or maximum < minimum
        ):
            raise ValueError(f"Borne maximale invalide pour {label!r}.")

        labels.add(label)
        validated.append(
            {"label": label.strip(), "min": minimum, "max": maximum}
        )

    validated.sort(key=lambda item: item["min"])

    previous_max = -1

    for index, band in enumerate(validated):
        minimum = band["min"]
        maximum = band["max"]

        if minimum <= previous_max:
            raise ValueError(
                f"Les tranches d'âge se chevauchent près de "
                f"{band['label']!r}."
            )

        if maximum is None and index != len(validated) - 1:
            raise ValueError(
                "Une tranche sans borne maximale doit être la dernière."
            )

        previous_max = math.inf if maximum is None else maximum

    return validated


def load_config(config_path: str | Path = "stat_config.json") -> dict[str, Any]:
    """Charge et valide le fichier JSON de configuration."""

    path = Path(config_path).expanduser().resolve()

    if not path.is_file():
        raise FileNotFoundError(
            f"Configuration introuvable : {path}"
        )

    try:
        with path.open("r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"JSON invalide à la ligne {exc.lineno}, "
            f"colonne {exc.colno}."
        ) from exc

    if not isinstance(config, dict):
        raise ValueError("La configuration doit être un objet JSON.")

    missing = [key for key in REQUIRED_CONFIG_KEYS if key not in config]

    if missing:
        raise ValueError(
            f"Paramètres manquants dans la configuration : "
            f"{', '.join(missing)}"
        )

    if not isinstance(config["columns"], dict) or not config["columns"]:
        raise ValueError("'columns' doit être un objet JSON non vide.")

    if (
        isinstance(config["start_row"], bool)
        or not isinstance(config["start_row"], int)
        or config["start_row"] < 1
    ):
        raise ValueError("'start_row' doit être un entier positif.")

    header_row = config.get("header_row", 1)

    if (
        isinstance(header_row, bool)
        or not isinstance(header_row, int)
        or header_row < 1
    ):
        raise ValueError("'header_row' doit être un entier positif.")

    if header_row >= config["start_row"]:
        raise ValueError(
            "'header_row' doit être inférieur à 'start_row'."
        )

    for key in ("input_file", "output_file", "sheet_name"):
        if not isinstance(config[key], str) or not config[key].strip():
            raise ValueError(f"'{key}' doit être une chaîne non vide.")

    config["columns"] = {
        name: _validate_excel_column(name, letter)
        for name, letter in config["columns"].items()
    }

    code_column = config.get("code_depart_column")

    if code_column is not None:
        config["code_depart_column"] = _validate_excel_column(
            "code_depart_column",
            code_column,
        )

    config["age_output_column"] = _validate_excel_column(
        "age_output_column",
        config.get("age_output_column", "I"),
    )

    age_header = config.get("age_header", "AGE")

    if not isinstance(age_header, str) or not age_header.strip():
        raise ValueError("'age_header' doit être une chaîne non vide.")

    config["age_header"] = age_header.strip()
    config["header_row"] = header_row

    try:
        config["_reference_date"] = date.fromisoformat(
            config["reference_date"]
        )
    except (TypeError, ValueError) as exc:
        raise ValueError(
            "'reference_date' doit respecter le format AAAA-MM-JJ."
        ) from exc

    date_formats = config.get(
        "date_formats",
        list(DEFAULT_DATE_FORMATS),
    )

    if (
        not isinstance(date_formats, list)
        or not date_formats
        or not all(
            isinstance(fmt, str) and fmt.strip()
            for fmt in date_formats
        )
    ):
        raise ValueError(
            "'date_formats' doit être une liste de formats de date."
        )

    config["date_formats"] = date_formats
    config["age_bands"] = _validate_age_bands(config["age_bands"])

    if "DATE NAISSANCE" not in config["columns"]:
        raise ValueError(
            "La colonne 'DATE NAISSANCE' est nécessaire pour calculer l'âge."
        )

    # Détecter les conflits de colonnes avant de modifier le fichier.
    configured_columns = {
        name: letter for name, letter in config["columns"].items()
    }

    age_column = config["age_output_column"]

    conflicting_fields = [
        name
        for name, letter in configured_columns.items()
        if letter == age_column and name != "AGE"
    ]

    if conflicting_fields:
        raise ValueError(
            f"La colonne AGE configurée ({age_column}) est déjà utilisée "
            f"par : {', '.join(conflicting_fields)}. "
            "Choisis une autre colonne pour 'age_output_column'."
        )

    return config


# ============================================================
# VALEURS ET DATES
# ============================================================

def is_empty(value: Any) -> bool:
    """Considère None, les chaînes vides et VIDE comme des valeurs absentes."""

    return value is None or (
        isinstance(value, str)
        and value.strip().casefold() in {"", "vide"}
    )


def parse_date(
    value: Any,
    formats: list[str] | tuple[str, ...] | None = None,
    excel_epoch=None,
) -> date | None:
    """Convertit une date Excel ou une chaîne en datetime.date."""

    if is_empty(value) or isinstance(value, bool):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        if not math.isfinite(value) or not 1000 <= value <= 100000:
            return None

        try:
            converted = (
                from_excel(value, excel_epoch)
                if excel_epoch is not None
                else from_excel(value)
            )
        except (OverflowError, TypeError, ValueError):
            return None

        if isinstance(converted, datetime):
            return converted.date()

        if isinstance(converted, date):
            return converted

        return None

    if isinstance(value, str):
        text = value.strip()

        for fmt in formats or DEFAULT_DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue

    return None


def calculate_age(
    birth_date: date | None,
    reference_date: date,
) -> int | None:
    """Calcule l'âge révolu en années complètes."""

    if birth_date is None or birth_date > reference_date:
        return None

    age = reference_date.year - birth_date.year

    if (reference_date.month, reference_date.day) < (
        birth_date.month,
        birth_date.day,
    ):
        age -= 1

    return age


def normalize_code(value: Any) -> str | None:
    """Normalise un code, y compris les valeurs ' 18 ' et '18.0'."""

    if is_empty(value):
        return None

    text = str(value).strip()

    try:
        number = Decimal(text)

        if number.is_finite() and number == number.to_integral_value():
            return str(int(number))
    except (InvalidOperation, ValueError):
        pass

    return text.casefold()


def classify_gender(civilite: Any) -> str:
    """Classe les civilités : Mme/Mlle = Femme, M/Monsieur = Homme."""

    if is_empty(civilite):
        return "VIDE"

    value = " ".join(str(civilite).strip().casefold().split())
    value = value.replace(".", "")

    if value in {"mme", "mlle", "mademoiselle"}:
        return "Femme"

    if value in {"m", "mr", "monsieur"}:
        return "Homme"

    return "VIDE"


def classify_age(age: int | None, age_bands: list[dict[str, Any]]) -> str:
    """Associe un âge à une tranche configurée."""

    if age is None:
        return "Âge inconnu"

    if age < 16:
        return "Moins de 16 ans"

    for band in age_bands:
        minimum = band["min"]
        maximum = band["max"]

        if age >= minimum and (maximum is None or age <= maximum):
            return band["label"]

    return "Âge inconnu"


# ============================================================
# CALCUL ET ÉCRITURE DE LA COLONNE AGE
# ============================================================

def add_age_column_to_input(config: dict[str, Any]) -> dict[str, Any]:
    """
    Ajoute ou actualise la colonne AGE dans le fichier source.

    Une sauvegarde .bak est créée avant la première modification.
    Les colonnes métier configurées ne sont jamais écrasées.
    """

    input_path = Path(config["input_file"]).expanduser().resolve()

    if not input_path.is_file():
        raise FileNotFoundError(
            f"Fichier source introuvable : {input_path}"
        )

    if input_path.suffix.lower() not in SUPPORTED_EXCEL_EXTENSIONS:
        raise ValueError(
            "Le fichier source doit être au format .xlsx ou .xlsm."
        )

    age_column = config["age_output_column"]
    age_index = column_index_from_string(age_column)
    header_row = config["header_row"]
    start_row = config["start_row"]

    # Ne pas autoriser la colonne AGE à écraser une colonne métier.
    for field_name, letter in config["columns"].items():
        if letter == age_column and field_name != "AGE":
            raise ValueError(
                f"La colonne {age_column} est déjà utilisée par "
                f"{field_name!r}. Choisis une autre colonne AGE."
            )

    keep_vba = input_path.suffix.lower() == ".xlsm"

    backup_path = input_path.with_suffix(input_path.suffix + ".bak")

    workbook = load_workbook(
        input_path,
        read_only=False,
        data_only=False,
        keep_vba=keep_vba,
    )

    try:
        sheet_name = config["sheet_name"]

        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Feuille {sheet_name!r} introuvable dans "
                f"{input_path.name}. Feuilles disponibles : "
                f"{', '.join(workbook.sheetnames)}"
            )

        worksheet = workbook[sheet_name]

        header_cell = worksheet.cell(
            row=header_row,
            column=age_index,
        )
        header_value = header_cell.value

        if (
            not is_empty(header_value)
            and str(header_value).strip().casefold()
            != config["age_header"].casefold()
        ):
            raise ValueError(
                f"La colonne {age_column} contient déjà l'en-tête "
                f"{header_value!r}. Vérifie 'age_output_column' dans "
                "stat_config.json avant de continuer."
            )

        birth_index = column_index_from_string(
            config["columns"]["DATE NAISSANCE"]
        )

        source_indexes = {
            column_index_from_string(letter)
            for letter in config["columns"].values()
        }

        # Repérer la dernière ligne comportant au moins une donnée métier.
        last_row = header_row

        for row_number in range(start_row, worksheet.max_row + 1):
            if any(
                not is_empty(worksheet.cell(row=row_number, column=index).value)
                for index in source_indexes
            ):
                last_row = row_number

        # Sauvegarde uniquement après validation, avant toute modification.
        if not backup_path.exists():
            shutil.copy2(input_path, backup_path)

        # Reprendre le style de la colonne précédente si disponible.
        source_header = worksheet.cell(
            row=header_row,
            column=max(1, age_index - 1),
        )

        if source_header.has_style:
            header_cell._style = copy(source_header._style)

        if source_header.alignment:
            header_cell.alignment = copy(source_header.alignment)
        else:
            header_cell.alignment = Alignment(horizontal="center")

        header_cell.value = config["age_header"]

        reference_date = config["_reference_date"]
        empty_value = config.get("empty_value", "VIDE")
        invalid_dates = 0
        processed_rows = 0

        for row_number in range(start_row, last_row + 1):
            birth_cell = worksheet.cell(
                row=row_number,
                column=birth_index,
            )
            age_cell = worksheet.cell(
                row=row_number,
                column=age_index,
            )

            birth_date = parse_date(
                birth_cell.value,
                config["date_formats"],
                workbook.epoch,
            )

            age = calculate_age(birth_date, reference_date)

            style_source = worksheet.cell(
                row=row_number,
                column=max(1, age_index - 1),
            )

            if style_source.has_style:
                age_cell._style = copy(style_source._style)

            if age is None:
                age_cell.value = empty_value

                # Une date absente n'est pas considérée invalide.
                if not is_empty(birth_cell.value):
                    invalid_dates += 1
            else:
                age_cell.value = age
                age_cell.number_format = "0"

            processed_rows += 1

        worksheet.column_dimensions[age_column].width = 12

        workbook.save(input_path)

        return {
            "path": input_path,
            "backup_path": backup_path,
            "rows": processed_rows,
            "invalid_dates": invalid_dates,
            "column": age_column,
        }

    finally:
        workbook.close()


# ============================================================
# PRÉPARATION DES EMPLOYÉS
# ============================================================

def _validate_input_file(config: dict[str, Any]) -> Path:
    """Vérifie l'existence et le format du fichier d'entrée."""

    input_path = Path(config["input_file"]).expanduser().resolve()

    if not input_path.is_file():
        raise FileNotFoundError(
            f"Fichier source introuvable : {input_path}"
        )

    if input_path.suffix.lower() not in SUPPORTED_EXCEL_EXTENSIONS:
        raise ValueError(
            "Le fichier source doit être au format .xlsx ou .xlsm."
        )

    output_path = Path(config["output_file"]).expanduser().resolve()

    if input_path == output_path:
        raise ValueError(
            "Le fichier source et le fichier de sortie doivent être différents."
        )

    return input_path


def _read_row_values(
    row: tuple[Any, ...],
    column_indexes: dict[str, int],
) -> dict[str, Any]:
    """Construit le dictionnaire des champs métier d'une ligne."""

    return {
        name: row[index - 1] if index - 1 < len(row) else None
        for name, index in column_indexes.items()
    }


def _row_is_empty(values: dict[str, Any], code_value: Any) -> bool:
    """Indique si aucune donnée métier ni aucun code de départ n'est présent."""

    return all(is_empty(value) for value in values.values()) and is_empty(
        code_value
    )


def _prepare_employee(
    row_number: int,
    values: dict[str, Any],
    config: dict[str, Any],
    excel_epoch,
) -> dict[str, Any]:
    """Prépare un employé et conserve les champs métier à la racine."""

    birth_date = parse_date(
        values.get("DATE NAISSANCE"),
        formats=config["date_formats"],
        excel_epoch=excel_epoch,
    )

    reference_date = config["_reference_date"]

    if birth_date is None or birth_date > reference_date:
        age = None
    else:
        age = calculate_age(birth_date, reference_date)

    employee = dict(values)

    # Compatibilité avec les modules existants.
    employee.update(
        {
            "source_row": row_number,
            "values": values,
            "gender": classify_gender(values.get("CIVILITE")),
            "birth_date": birth_date,
            "age": age,
            "age_band": classify_age(age, config["age_bands"]),
        }
    )

    return employee


def load_employees(config: dict[str, Any]) -> dict[str, Any]:
    """
    Charge les employés depuis Excel.

    Chaque employé conserve :
    - ses colonnes métier à la racine (GRADE, CIVILITE, etc.) ;
    - le dictionnaire historique 'values' pour compatibilité ;
    - son âge, son sexe et sa tranche d'âge.

    Si 'code_depart_column' est configurée, les lignes avec le code 18
    sont exclues avant les statistiques.
    """

    input_path = _validate_input_file(config)
    columns = config["columns"]

    missing_columns = REQUIRED_EMPLOYEE_COLUMNS - set(columns)

    if missing_columns:
        raise ValueError(
            "Colonnes nécessaires manquantes : "
            + ", ".join(sorted(missing_columns))
        )

    column_indexes = {
        name: column_index_from_string(letter)
        for name, letter in columns.items()
    }

    code_depart_column = config.get("code_depart_column")
    code_depart_index = (
        column_index_from_string(code_depart_column)
        if code_depart_column
        else None
    )

    max_column = max(
        [*column_indexes.values(), code_depart_index or 0]
    )

    keep_vba = input_path.suffix.lower() == ".xlsm"

    workbook = load_workbook(
        input_path,
        read_only=True,
        data_only=True,
        keep_vba=keep_vba,
    )

    try:
        sheet_name = config["sheet_name"]

        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Feuille {sheet_name!r} introuvable. "
                f"Feuilles disponibles : {', '.join(workbook.sheetnames)}"
            )

        worksheet = workbook[sheet_name]

        employees = []
        invalid_birth_dates = 0
        future_birth_dates = 0
        code_depart_filtered_count = 0
        empty_rows_skipped = 0

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=config["start_row"],
                max_col=max_column,
                values_only=True,
            ),
            start=config["start_row"],
        ):
            values = _read_row_values(row, column_indexes)

            code_value = (
                row[code_depart_index - 1]
                if code_depart_index is not None
                and code_depart_index - 1 < len(row)
                else None
            )

            if _row_is_empty(values, code_value):
                empty_rows_skipped += 1
                continue

            if (
                code_depart_index is not None
                and normalize_code(code_value) == "18"
            ):
                code_depart_filtered_count += 1
                continue

            employee = _prepare_employee(
                row_number=row_number,
                values=values,
                config=config,
                excel_epoch=workbook.epoch,
            )

            birth_date = employee["birth_date"]

            if birth_date is None:
                if not is_empty(values.get("DATE NAISSANCE")):
                    invalid_birth_dates += 1
            elif birth_date > config["_reference_date"]:
                future_birth_dates += 1

            employees.append(employee)

        return {
            "employees": employees,
            "invalid_birth_dates": invalid_birth_dates,
            "future_birth_dates": future_birth_dates,
            "code_depart_filtered_count": code_depart_filtered_count,
            "empty_rows_skipped": empty_rows_skipped,
            "config": config,
        }

    finally:
        workbook.close()
