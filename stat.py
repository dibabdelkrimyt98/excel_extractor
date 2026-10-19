
"""Point d'entrée principal des statistiques Excel Extractor."""

import sys
import time
from pathlib import Path
from typing import Any, Callable, TypeVar

from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

from stat_utils import (
    add_age_column_to_input,
    load_config,
    load_employees,
)
from stat_sexe import calculate_gender_statistics
from stat_age import calculate_age_statistics
from stat_effectifs import calculate_effectifs_statistics
from stat_writer import write_statistics


T = TypeVar("T")


# ============================================================
# MESURE DU TEMPS
# ============================================================

def timed_operation(
    label: str,
    operation: Callable[[], T],
) -> T:
    """Exécute une opération et affiche son temps de traitement."""

    start = time.perf_counter()

    try:
        return operation()
    finally:
        elapsed = time.perf_counter() - start
        print(
            f"[TEMPS] {label} : {elapsed:.3f} seconde(s)",
            flush=True,
        )


# ============================================================
# STYLES EXCEL
# ============================================================

def get_excel_styles() -> dict[str, Any]:
    """Retourne les styles communs aux tableaux de statistiques."""

    thin_gray = Side(style="thin", color="D9E2F3")

    return {
        "title_font": Font(bold=True, size=12),
        "header_font": Font(bold=True, color="FFFFFF"),
        "header_fill": PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        ),
        "label_font": Font(bold=True),
        "border": Border(
            left=thin_gray,
            right=thin_gray,
            top=thin_gray,
            bottom=thin_gray,
        ),
    }


def style_header_row(ws, row: int, start_col: int, headers: list[str]):
    """Écrit et met en forme une ligne d'en-têtes."""

    styles = get_excel_styles()

    for offset, header in enumerate(headers):
        cell = ws.cell(
            row=row,
            column=start_col + offset,
            value=header,
        )
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.border = styles["border"]
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def style_data_row(ws, row: int, start_col: int, end_col: int):
    """Applique les bordures et l'alignement à une ligne."""

    styles = get_excel_styles()

    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = styles["border"]
        cell.alignment = Alignment(
            vertical="center",
            horizontal="left" if col == start_col else "center",
            wrap_text=True,
        )


# ============================================================
# MOYENNES D'ÂGE
# ============================================================

def append_average_ages(
    output_file: str | Path,
    age_stats: dict[str, Any],
) -> None:
    """Ajoute les moyennes d'âge à la fin de Tranches d'âge."""

    output_path = Path(output_file).expanduser().resolve()

    if not output_path.is_file():
        raise FileNotFoundError(
            f"Fichier de statistiques introuvable : {output_path}"
        )

    wb = None

    try:
        wb = load_workbook(output_path)

        sheet_name = "Tranches d'âge"

        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"La feuille {sheet_name!r} est absente. "
                f"Feuilles disponibles : {', '.join(wb.sheetnames)}"
            )

        ws = wb[sheet_name]
        styles = get_excel_styles()

        start_row = ws.max_row + 2

        title = ws.cell(
            row=start_row,
            column=1,
            value="RÉCAPITULATIF DES MOYENNES D'ÂGE",
        )
        title.font = styles["title_font"]

        ws.merge_cells(
            start_row=start_row,
            start_column=1,
            end_row=start_row,
            end_column=3,
        )

        header_row = start_row + 1

        style_header_row(
            ws,
            header_row,
            1,
            [
                "Population",
                "Moyenne d'âge (ans)",
                "Nombre d'âges valides",
            ],
        )

        averages = age_stats.get("average_ages", {})
        counts = age_stats.get("average_age_counts", {})

        for offset, population in enumerate(
            ("Femme", "Homme", "Total"),
            start=1,
        ):
            row = header_row + offset
            average = averages.get(population)

            ws.cell(row=row, column=1, value=population)
            ws.cell(
                row=row,
                column=2,
                value=round(average, 2) if average is not None else "N/A",
            )
            ws.cell(
                row=row,
                column=3,
                value=counts.get(population, 0),
            )

            ws.cell(row=row, column=1).font = styles["label_font"]
            style_data_row(ws, row, 1, 3)

            if average is not None:
                ws.cell(row=row, column=2).number_format = "0.00"

        ws.column_dimensions["A"].width = max(
            ws.column_dimensions["A"].width or 0,
            22,
        )
        ws.column_dimensions["B"].width = max(
            ws.column_dimensions["B"].width or 0,
            24,
        )
        ws.column_dimensions["C"].width = max(
            ws.column_dimensions["C"].width or 0,
            25,
        )

        wb.save(output_path)

        print(
            f"[EXCEL] Moyennes d'âge ajoutées dans '{sheet_name}'.",
            flush=True,
        )

    finally:
        if wb is not None:
            wb.close()


# ============================================================
# STATISTIQUES D'EFFECTIFS
# ============================================================


def append_effectifs_statistics(
    output_file: str | Path,
    effectifs_stats: dict[str, Any],
) -> None:
    """
    Crée ou remplace les feuilles Effectifs et Contrôle motifs.

    Contenu :
    - effectifs par type de contrat (CDD et CDI, CCTA inclus dans CDI) ;
    - répartition des motifs de départ ;
    - grades par sexe ;
    - contrôles de qualité des données ;
    - liste des lignes dont le motif de départ est vide.
    """

    output_path = Path(output_file).expanduser().resolve()

    if not output_path.is_file():
        raise FileNotFoundError(
            f"Fichier de statistiques introuvable : {output_path}"
        )

    wb = None

    try:
        wb = load_workbook(output_path)

        # ====================================================
        # FEUILLE 1 : EFFECTIFS
        # ====================================================

        sheet_name = "Effectifs"

        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

        ws = wb.create_sheet(sheet_name)
        styles = get_excel_styles()
        ws.sheet_view.showGridLines = False

        # ----------------------------------------------------
        # TABLEAU 1 : TYPES DE CONTRAT
        # ----------------------------------------------------

        row = 1

        ws.cell(
            row=row,
            column=1,
            value="1. RÉPARTITION PAR TYPE DE CONTRAT",
        ).font = styles["title_font"]

        row += 1

        style_header_row(
            ws,
            row,
            1,
            ["Type de contrat", "Nombre d'employés"],
        )

        contracts = effectifs_stats["contracts"]
        contract_counts = contracts["counts"]

        # CCTA est regroupé dans CDI par stat_effectifs.py.
        for contract in ("CDD", "CDI"):
            row += 1

            ws.cell(row=row, column=1, value=contract)
            ws.cell(
                row=row,
                column=2,
                value=contract_counts.get(contract, 0),
            )

            style_data_row(ws, row, 1, 2)

        row += 1
        ws.cell(
            row=row,
            column=1,
            value="Total des contrats connus",
        )
        ws.cell(
            row=row,
            column=2,
            value=contracts["total_known"],
        )
        style_data_row(ws, row, 1, 2)

        row += 1
        ws.cell(
            row=row,
            column=1,
            value="Contrats vides ou VIDE",
        )
        ws.cell(
            row=row,
            column=2,
            value=contracts["empty_count"],
        )
        style_data_row(ws, row, 1, 2)

        row += 1
        ws.cell(
            row=row,
            column=1,
            value="Autres valeurs de contrat",
        )
        ws.cell(
            row=row,
            column=2,
            value=contracts["unknown_count"],
        )
        style_data_row(ws, row, 1, 2)

        # ----------------------------------------------------
        # TABLEAU 2 : MOTIFS DE DÉPART
        # ----------------------------------------------------

        row += 2

        ws.cell(
            row=row,
            column=1,
            value="2. RÉPARTITION PAR MOTIF DE DÉPART",
        ).font = styles["title_font"]

        row += 1

        style_header_row(
            ws,
            row,
            1,
            ["Motif de départ", "Nombre d'occurrences"],
        )

        departure_stats = effectifs_stats["departure_reasons"]
        reasons = departure_stats["reasons"]

        for reason in reasons:
            row += 1

            ws.cell(
                row=row,
                column=1,
                value=reason["motif"],
            )
            ws.cell(
                row=row,
                column=2,
                value=reason["count"],
            )

            style_data_row(ws, row, 1, 2)

        row += 1
        ws.cell(
            row=row,
            column=1,
            value="Total des motifs comptabilisés",
        )
        ws.cell(
            row=row,
            column=2,
            value=departure_stats["total_counted"],
        )
        style_data_row(ws, row, 1, 2)

        row += 1
        ws.cell(
            row=row,
            column=1,
            value="Motifs vides ou VIDE ignorés",
        )
        ws.cell(
            row=row,
            column=2,
            value=departure_stats["ignored_count"],
        )
        style_data_row(ws, row, 1, 2)

        # ----------------------------------------------------
        # TABLEAU 3 : GRADES PAR SEXE
        # ----------------------------------------------------

        row += 2

        ws.cell(
            row=row,
            column=1,
            value="3. RÉPARTITION DES GRADES PAR SEXE",
        ).font = styles["title_font"]

        row += 1

        grade_names = (
            "Encadrement",
            "Cadre",
            "Maitrise",
            "Execution",
        )

        style_header_row(
            ws,
            row,
            1,
            [
                "Sexe",
                *grade_names,
                "Total",
            ],
        )

        grades_stats = effectifs_stats["grades"]["counts"]

        for gender in ("Femme", "Homme"):
            row += 1
            gender_counts = grades_stats[gender]

            ws.cell(
                row=row,
                column=1,
                value=gender,
            )

            for col, grade in enumerate(grade_names, start=2):
                ws.cell(
                    row=row,
                    column=col,
                    value=gender_counts.get(grade, 0),
                )

            ws.cell(
                row=row,
                column=6,
                value=sum(gender_counts.values()),
            )

            style_data_row(ws, row, 1, 6)

        # Total femmes + hommes pour chaque grade.
        row += 1
        ws.cell(row=row, column=1, value="Total")

        for col, grade in enumerate(grade_names, start=2):
            total_grade = sum(
                grades_stats[gender].get(grade, 0)
                for gender in ("Femme", "Homme")
            )

            ws.cell(
                row=row,
                column=col,
                value=total_grade,
            )

        ws.cell(
            row=row,
            column=6,
            value=sum(
                sum(grades_stats[gender].values())
                for gender in ("Femme", "Homme")
            ),
        )

        style_data_row(ws, row, 1, 6)

        # ----------------------------------------------------
        # INFORMATIONS DE CONTRÔLE
        # ----------------------------------------------------

        row += 2

        ws.cell(
            row=row,
            column=1,
            value="Contrôles de qualité des données",
        ).font = styles["title_font"]

        quality_checks = [
            (
                "Grades non reconnus",
                effectifs_stats["grades"]["unknown_grade_count"],
            ),
            (
                "Grades vides ou VIDE",
                effectifs_stats["grades"]["empty_grade_count"],
            ),
            (
                "Grades avec civilité non reconnue",
                effectifs_stats["grades"]["unknown_gender_count"],
            ),
        ]

        for label, value in quality_checks:
            row += 1

            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)

            style_data_row(ws, row, 1, 2)

        # Largeur des colonnes.
        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 16

        ws.freeze_panes = "A3"

        # ====================================================
        # FEUILLE 2 : CONTRÔLE DES MOTIFS VIDES
        # ====================================================

        audit_sheet_name = "Contrôle motifs"

        if audit_sheet_name in wb.sheetnames:
            del wb[audit_sheet_name]

        audit_ws = wb.create_sheet(audit_sheet_name)
        audit_ws.sheet_view.showGridLines = False

        audit_headers = [
            "Ligne source",
            "CCP",
            "Motif brut",
            "Anomalie",
        ]

        audit_ws.append(audit_headers)

        style_header_row(
            audit_ws,
            1,
            1,
            audit_headers,
        )

        unread_rows = departure_stats.get("unread_rows", [])

        for item in unread_rows:
            audit_ws.append([
                item.get("ligne_source", "Inconnue"),
                item.get("CCP", ""),
                item.get("motif_brut", ""),
                item.get(
                    "anomalie",
                    "Motif vide ou non récupéré",
                ),
            ])

            current_row = audit_ws.max_row

            style_data_row(
                audit_ws,
                current_row,
                1,
                4,
            )

        audit_ws.freeze_panes = "A2"
        audit_ws.auto_filter.ref = audit_ws.dimensions

        audit_ws.column_dimensions["A"].width = 16
        audit_ws.column_dimensions["B"].width = 22
        audit_ws.column_dimensions["C"].width = 45
        audit_ws.column_dimensions["D"].width = 35

        # ====================================================
        # SAUVEGARDE
        # ====================================================

        wb.save(output_path)

        print(
            f"[EXCEL] Feuille '{sheet_name}' écrite.",
            flush=True,
        )
        print(
            f"[EXCEL] Feuille '{audit_sheet_name}' écrite : "
            f"{len(unread_rows)} ligne(s) à vérifier.",
            flush=True,
        )

    finally:
        if wb is not None:
            wb.close()
            

# ============================================================
# PROGRAMME PRINCIPAL
# ============================================================

def main() -> int:
    """Exécute toutes les statistiques et mesure leur durée."""

    total_start = time.perf_counter()

    try:
        # 1. Chargement de la configuration.
        config = timed_operation(
            "Chargement de stat_config.json",
            lambda: load_config("stat_config.json"),
        )

        # 2. Calcul de la colonne AGE dans le fichier source.
        age_column_result = timed_operation(
            "Calcul et mise à jour de la colonne AGE",
            lambda: add_age_column_to_input(config),
        )

        print("\n=== Calcul de la colonne AGE ===")
        print(f"Colonne mise à jour : {age_column_result['column']}")
        print(f"Lignes traitées    : {age_column_result['rows']}")
        print(
            "Dates absentes/invalides : "
            f"{age_column_result['invalid_dates']}"
        )
        print(f"Sauvegarde initiale : {age_column_result['backup_path']}")

        # 3. Chargement des données.
        employees_data = timed_operation(
            "Chargement des données employés",
            lambda: load_employees(config),
        )

        # 4. Statistiques par sexe.
        gender_stats = timed_operation(
            "Calcul des statistiques par sexe",
            lambda: calculate_gender_statistics(employees_data),
        )

        # 5. Statistiques d'âge.
        age_stats = timed_operation(
            "Calcul des statistiques d'âge et des moyennes",
            lambda: calculate_age_statistics(employees_data),
        )

        # 6. Statistiques d'effectifs.
        effectifs_stats = timed_operation(
            "Calcul des statistiques d'effectifs",
            lambda: calculate_effectifs_statistics(employees_data),
        )

        # 7. Génération du classeur principal.
        output_path = timed_operation(
            "Génération du fichier principal de statistiques",
            lambda: write_statistics(
                gender_stats=gender_stats,
                age_stats=age_stats,
                output_file=config["output_file"],
                employees_data=employees_data,
            ),
        )

        # Si le writer ne renvoie pas le chemin, on utilise celui de config.
        if output_path is None:
            output_path = Path(config["output_file"]).expanduser().resolve()

        # 8. Ajout des moyennes d'âge.
        timed_operation(
            "Écriture des moyennes dans Tranches d'âge",
            lambda: append_average_ages(
                output_file=output_path,
                age_stats=age_stats,
            ),
        )

        # 9. Ajout des statistiques d'effectifs.
        timed_operation(
            "Écriture de la feuille Effectifs",
            lambda: append_effectifs_statistics(
                output_file=output_path,
                effectifs_stats=effectifs_stats,
            ),
        )

        # ----------------------------------------------------
        # BILAN FINAL
        # ----------------------------------------------------

        print("\n" + "=" * 65)
        print("       EXCEL EXTRACTOR - STATISTIQUES")
        print("=" * 65)

        print(f"Employés analysés : {gender_stats['total']}")
        print(f"Femmes            : {gender_stats['counts']['Femme']}")
        print(f"Hommes            : {gender_stats['counts']['Homme']}")
        print(
            "Civilités inconnues : "
            f"{gender_stats['counts']['VIDE']}"
        )

        print("\n--- Types de contrat ---")

        for contract, count in effectifs_stats["contracts"]["counts"].items():
            print(f"{contract:<10}: {count}")

        print(
            "Contrats vides ou VIDE : "
            f"{effectifs_stats['contracts']['empty_count']}"
        )
        print(
            "Autres valeurs de contrat : "
            f"{effectifs_stats['contracts']['unknown_count']}"
        )

        print("\n--- Motifs de départ ---")

        for reason in effectifs_stats["departure_reasons"]["reasons"]:
            print(f"{reason['motif']} : {reason['count']}")

        print(
            "Motifs vides ou VIDE ignorés : "
            f"{effectifs_stats['departure_reasons']['ignored_count']}"
        )

        print("\n--- Grades par sexe ---")

        for gender in ("Femme", "Homme"):
            print(f"\n{gender} :")

            for grade, count in effectifs_stats["grades"]["counts"][
                gender
            ].items():
                print(f"  {grade:<15}: {count}")

        print("\n--- Statistiques d'âge ---")

        print(
            "Lignes exclues (Code Depart = 18) : "
            f"{age_stats['code_depart_filtered_count']}"
        )
        print(
            "Dates de naissance absentes/invalides : "
            f"{age_stats['invalid_birth_dates']}"
        )
        print(
            "Dates de naissance futures : "
            f"{age_stats['future_birth_dates']}"
        )
        print(
            "Lignes entièrement vides ignorées : "
            f"{age_stats['empty_rows_skipped']}"
        )
        print(f"Date de référence : {config['reference_date']}")

        averages = age_stats["average_ages"]
        average_counts = age_stats["average_age_counts"]

        print("\n--- Moyennes d'âge ---")

        for population in ("Femme", "Homme", "Total"):
            average = averages[population]
            count = average_counts[population]

            if average is None:
                print(f"{population:<10}: N/A (aucun âge valide)")
            else:
                print(
                    f"{population:<10}: {average:.2f} ans "
                    f"({count} âge(s) valide(s))"
                )

        print("\n--- Fichier généré ---")
        print(f"Fichier : {output_path}")

        print("\n--- Temps de traitement ---")
        print(
            f"Temps total : "
            f"{time.perf_counter() - total_start:.3f} seconde(s)"
        )

        print("=" * 65)

        return 0

    except (
        FileNotFoundError,
        PermissionError,
        ValueError,
        OSError,
    ) as exc:
        print(f"\nERREUR : {exc}", file=sys.stderr)
        return 1

    except Exception as exc:
        print(
            f"\nERREUR INATTENDUE : "
            f"{type(exc).__name__}: {exc}",
            file=sys.stderr,
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
